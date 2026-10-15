
#! STEP 1: Import the libraries
from openai import OpenAI
from dotenv import load_dotenv
import os
import json
from openpyxl import Workbook, load_workbook


#! STEP 2: Connect to OpenAI
load_dotenv()

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)

print("OpenAI Connected")

FILE_PATH = "./MY_WORK/Project/10_Excel_Analyzer_CRUD/test.xlsx"

print("Excel file path:", os.path.abspath(FILE_PATH))


#! STEP 3: Excel functions

# Note:
# - The ID is always the FIRST column.
# - Errors are caught in the chat loop.


def create_excel(headers, rows):
    # Validate headers
    if not headers:
        return "Error: At least one column header is required."

    if any(not isinstance(h, str) or not h.strip() for h in headers):
        return "Error: Column headers must be non-empty strings."

    if len(set(h.lower() for h in headers)) != len(headers):
        return "Error: Duplicate column headers are not allowed."

    if any(len(row) != len(headers) for row in rows):
        return "Error: Every row must have the same number of values as headers."

    os.makedirs(os.path.dirname(FILE_PATH), exist_ok=True)

    wb = Workbook()
    sheet = wb.active

    sheet.append(headers)

    for row in rows:
        sheet.append(row)

    wb.save(FILE_PATH)

    return "Excel file created successfully."


def read_excel():
    if not os.path.exists(FILE_PATH):
        return "File does not exist."

    wb = load_workbook(FILE_PATH, data_only=True)
    sheet = wb.active

    data = list(sheet.values)

    return json.dumps(data, default=str, ensure_ascii=False)


def add_column(column_name):
    if not os.path.exists(FILE_PATH):
        return "Error: Excel file does not exist."

    if not column_name or not column_name.strip():
        return "Error: Column name cannot be empty."

    wb = load_workbook(FILE_PATH)
    sheet = wb.active

    headers = [
        str(cell.value).lower()
        for cell in sheet[1]
    ]

    # Do not add the same column twice.
    if column_name.lower() in headers:
        return "Column already exists."

    # Add the new column at the end.
    sheet.cell(
        row=1,
        column=sheet.max_column + 1,
        value=column_name
    )

    wb.save(FILE_PATH)

    return f"Column '{column_name}' added."


def delete_column(column_name):
    if not os.path.exists(FILE_PATH):
        return "Error: Excel file does not exist."

    if not column_name or not column_name.strip():
        return "Error: Column name cannot be empty."

    wb = load_workbook(FILE_PATH)
    sheet = wb.active

    headers = [
        str(cell.value).lower()
        for cell in sheet[1]
    ]

    # Find the requested column.
    if column_name.lower() not in headers:
        return f"Error: Column '{column_name}' not found."

    col = headers.index(column_name.lower()) + 1

    # Protect the first column (ID).
    if col == 1:
        return "Error: The ID column (first column) cannot be deleted."

    # Delete the column header and all data in that column.
    sheet.delete_cols(col, 1)

    wb.save(FILE_PATH)

    return (
        f"Column '{column_name}' and all its data "
        f"deleted successfully."
    )


def fill_column(column_name, values):
    if not os.path.exists(FILE_PATH):
        return "Error: Excel file does not exist."

    wb = load_workbook(FILE_PATH)
    sheet = wb.active

    headers = [
        str(cell.value).lower()
        for cell in sheet[1]
    ]

    if column_name.lower() not in headers:
        return "Column not found."

    col = headers.index(column_name.lower()) + 1

    # Exactly one value per data row.
    total_rows = sheet.max_row - 1

    if len(values) != total_rows:
        return (
            f"Error: {total_rows} rows exist but "
            f"{len(values)} values given."
        )

    # Data starts from row 2.
    for i, value in enumerate(values):
        sheet.cell(
            row=i + 2,
            column=col,
            value=value
        )

    wb.save(FILE_PATH)

    return f"Column '{column_name}' filled for all rows."


def insert_row(data):
    if not os.path.exists(FILE_PATH):
        return "Error: Excel file does not exist."

    wb = load_workbook(FILE_PATH)
    sheet = wb.active

    headers = [
        str(cell.value).lower()
        for cell in sheet[1]
    ]

    if not headers or sheet.cell(1, 1).value is None:
        return "Error: Excel file has no valid headers."

    # Normalize user-provided keys.
    new_data = {}

    for key, value in data.items():
        normalized_key = key.lower()

        if normalized_key == "id":
            normalized_key = headers[0]

        if normalized_key not in headers:
            return (
                f"Error: Column '{key}' not found. "
                f"Columns: {headers}"
            )

        new_data[normalized_key] = value

    id_column = headers[0]

    # Automatically generate an ID if one was not provided.
    if new_data.get(id_column) in (None, ""):
        biggest = 0

        for r in range(2, sheet.max_row + 1):
            value = sheet.cell(row=r, column=1).value

            if isinstance(value, int) and not isinstance(value, bool):
                if value > biggest:
                    biggest = value

        new_data[id_column] = biggest + 1

    else:
        # Prevent duplicate IDs.
        for r in range(2, sheet.max_row + 1):
            existing_id = sheet.cell(row=r, column=1).value

            if str(existing_id) == str(new_data[id_column]):
                return (
                    f"Error: ID {new_data[id_column]} "
                    f"already exists."
                )

    # Build row in the same order as the headers.
    new_row = [
        new_data.get(header)
        for header in headers
    ]

    sheet.append(new_row)
    wb.save(FILE_PATH)

    return f"Row inserted with ID {new_data[id_column]}."


def update_row(id, data):
    # An ID is required for updates.
    if id is None or str(id).strip() == "":
        return "ID is required for update. Ask the user to provide the ID."

    if not data:
        return "No data to update. Ask the user which values to change."

    if not os.path.exists(FILE_PATH):
        return "Error: Excel file does not exist."

    wb = load_workbook(FILE_PATH)
    sheet = wb.active

    headers = [
        str(cell.value).lower()
        for cell in sheet[1]
    ]

    # Validate all requested columns before changing anything.
    for key in data:
        if key.lower() not in headers:
            return f"Error: column '{key}' not found. Columns: {headers}"

    # Find the record by ID.
    for r in range(2, sheet.max_row + 1):
        existing_id = sheet.cell(row=r, column=1).value

        if str(existing_id) == str(id):
            for key, value in data.items():
                col = headers.index(key.lower()) + 1

                sheet.cell(
                    row=r,
                    column=col,
                    value=value
                )

            wb.save(FILE_PATH)

            return f"Record with ID {id} updated."

    return f"No record found with ID {id}."


def delete_row(id):
    # An ID is required for deletion.
    if id is None or str(id).strip() == "":
        return "ID is required for delete. Ask the user to provide the ID."

    if not os.path.exists(FILE_PATH):
        return "Error: Excel file does not exist."

    wb = load_workbook(FILE_PATH)
    sheet = wb.active

    # Find the record by ID.
    for r in range(2, sheet.max_row + 1):
        existing_id = sheet.cell(row=r, column=1).value

        if str(existing_id) == str(id):
            sheet.delete_rows(r, 1)

            wb.save(FILE_PATH)

            return f"Record with ID {id} deleted."

    return f"No record found with ID {id}."


#! STEP 4: Define AI tools

tools = [
    {
        "type": "function",
        "function": {
            "name": "create_excel",
            "description": "Create an Excel file with headers and rows.",
            "parameters": {
                "type": "object",
                "properties": {
                    "headers": {
                        "type": "array",
                        "items": {"type": "string"}
                    },
                    "rows": {
                        "type": "array",
                        "items": {
                            "type": "array",
                            "items": {
                                "type": ["number", "string", "boolean", "null"]
                            }
                        }
                    }
                },
                "required": ["headers", "rows"],
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "read_excel",
            "description": "Read all content of the Excel file.",
            "parameters": {
                "type": "object",
                "properties": {},
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "add_column",
            "description": "Add a new empty column to the Excel file.",
            "parameters": {
                "type": "object",
                "properties": {
                    "column_name": {"type": "string"}
                },
                "required": ["column_name"],
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "delete_column",
            "description": (
                "Delete a column by name, including its header "
                "and all its data. Never delete the first ID column."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "column_name": {"type": "string"}
                },
                "required": ["column_name"],
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "fill_column",
            "description": (
                "Fill a column for ALL data rows. "
                "Provide exactly one value per data row, in order."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "column_name": {"type": "string"},
                    "values": {
                        "type": "array",
                        "items": {
                            "type": ["number", "string", "boolean", "null"]
                        }
                    }
                },
                "required": ["column_name", "values"],
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "insert_row",
            "description": (
                "Insert one new row. Keys are column names. "
                "ID is optional and will be generated if omitted."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "data": {
                        "type": "object",
                        "additionalProperties": {
                            "type": ["string", "number", "boolean", "null"]
                        }
                    }
                },
                "required": ["data"],
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "update_row",
            "description": (
                "Update a row by ID. Never guess an ID. "
                "If the user did not provide an ID, ask for it."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "id": {"type": ["string", "number"]},
                    "data": {
                        "type": "object",
                        "additionalProperties": {
                            "type": ["string", "number", "boolean", "null"]
                        }
                    }
                },
                "required": ["data"],
                "additionalProperties": False
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "delete_row",
            "description": (
                "Delete a row by ID. Never guess an ID. "
                "If the user did not provide an ID, ask for it."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "id": {"type": ["string", "number"]}
                },
                "required": ["id"],
                "additionalProperties": False
            }
        }
    }
]


#! STEP 5: System message

messages = [
    {
        "role": "system",
        "content": """
You are an Excel file agent working on test.xlsx.

RULES:

GENERAL:
- For any question about existing data, ALWAYS call read_excel first.
- Never guess or invent data.
- If no record is found, clearly say no matching record was found.
- If a tool returns an error, show the error message to the user.
- After any change, briefly tell the user what happened.

CREATE:
- Use create_excel when the user asks to create or replace the Excel file.
- Make sure the user has clearly requested creation or replacement.

READ:
- Read the Excel file before answering questions about existing records,
  values, column names, counts, or other existing data.

ADD COLUMN:
- Call add_column.
- The new column remains empty.

DELETE COLUMN:
- Call delete_column with the requested column name.
- The column header and every cell's data in that column will be deleted.
- Never delete the first column, which contains IDs.
- If the column name is unclear, ask the user.
- If the requested column is the ID column, explain that it cannot be deleted.

FILL COLUMN:
- First call read_excel to count the data rows.
- Then call fill_column with exactly one value per data row.
- Do not invent values unless the user asks for sample/generated values.

INSERT:
- Call insert_row.
- If the user does not provide an ID, do not ask for one.
- The function generates an ID automatically.
- Never invent a value for another field that the user has not specified,
  unless the user asks you to generate sample data.

UPDATE:
- An ID is REQUIRED.
- If the user did not give an ID, ask for it and do not call update_row.
- "id=14", "ID 14", or "employee id 14" means ID 14.
- Do not ask the user to confirm an ID they explicitly provided.
- Only update the fields the user requested.

DELETE ROW:
- An ID is REQUIRED.
- If the user did not give an ID, ask for it and do not call delete_row.
- Never guess an ID.
- "id=14", "ID 14", or "employee id 14" means ID 14.

- Do not claim an operation succeeded unless the tool result confirms it.
"""
    }
]


#! STEP 6: Chat loop

while True:
    user_input = input("\nYou: ").strip()

    if user_input.lower() in ["exit", "quit"]:
        print("Agent: Bye!")
        break

    if not user_input:
        continue

    messages.append({
        "role": "user",
        "content": user_input
    })

    # The AI may call several tools before answering.
    while True:
        response = client.chat.completions.create(
            model="gpt-4.1-mini",
            messages=messages,
            tools=tools,
            tool_choice="auto"
        )

        msg = response.choices[0].message

        # Preserve the assistant message, including any tool calls.
        messages.append(msg)

        # If no tool is needed, print the answer.
        if not msg.tool_calls:
            print("\nAgent:", msg.content or "")
            break

        # STEP 7: Execute the requested tools.
        for tool_call in msg.tool_calls:
            name = tool_call.function.name

            try:
                args = json.loads(
                    tool_call.function.arguments
                )

                if name == "create_excel":
                    result = create_excel(
                        args["headers"],
                        args["rows"]
                    )

                elif name == "read_excel":
                    result = read_excel()

                elif name == "add_column":
                    result = add_column(
                        args["column_name"]
                    )

                elif name == "delete_column":
                    result = delete_column(
                        args["column_name"]
                    )

                elif name == "fill_column":
                    result = fill_column(
                        args["column_name"],
                        args["values"]
                    )

                elif name == "insert_row":
                    result = insert_row(
                        args["data"]
                    )

                elif name == "update_row":
                    result = update_row(
                        args.get("id"),
                        args.get("data")
                    )

                elif name == "delete_row":
                    result = delete_row(
                        args.get("id")
                    )

                else:
                    result = "Error: Tool not recognized."

            except Exception as ex:
                result = f"Error: {str(ex)}"

            # Send the tool result back to the AI.
            messages.append({
                "role": "tool",
                "tool_call_id": tool_call.id,
                "content": str(result)
            })
